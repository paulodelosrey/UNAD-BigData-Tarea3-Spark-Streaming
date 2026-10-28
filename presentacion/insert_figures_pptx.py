import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

script_dir = os.path.dirname(os.path.abspath(__file__))
base_dir = os.path.dirname(script_dir)
pptx_path = os.path.join(script_dir, "Sustentación Tarea 3 Big Data UNAD.pptx")
backup_path = os.path.join(script_dir, "Sustentación Tarea 3 Big Data UNAD_backup.pptx")
img_dir = os.path.join(base_dir, "output")

# Backup
if not os.path.exists(backup_path):
    shutil.copyfile(pptx_path, backup_path)
    print("Backup creado en:", backup_path)

prs = Presentation(pptx_path)

# -------------------------------------------------------------
# SLIDE 4: Agregar Figura 1 (Comparativo PM) y Figura 3 (Top Estaciones)
# -------------------------------------------------------------
slide4 = prs.slides[3]

fig1_path = os.path.join(img_dir, "figura1_pm_departamentos.png")
fig3_path = os.path.join(img_dir, "figura3_top_estaciones_criticas.png")

if os.path.exists(fig1_path):
    # Left image: Figura 1
    slide4.shapes.add_picture(fig1_path, Inches(0.67), Inches(2.70), width=Inches(5.85), height=Inches(4.40))
    print("Figura 1 agregada a Diapositiva 4")

if os.path.exists(fig3_path):
    # Right image: Figura 3
    slide4.shapes.add_picture(fig3_path, Inches(6.80), Inches(2.70), width=Inches(5.85), height=Inches(4.40))
    print("Figura 3 agregada a Diapositiva 4")


# -------------------------------------------------------------
# SLIDE 5: Reemplazar placeholders y agregar Figura 2 y Figura 4
# -------------------------------------------------------------
slide5 = prs.slides[4]

fig2_path = os.path.join(img_dir, "figura2_distribucion_ica.png")
fig4_path = os.path.join(img_dir, "figura4_dispersion_clima_contaminacion.png")

# Limpiar texto de placeholder
for shape in slide5.shapes:
    if shape.has_text_frame:
        if "[ESPACIO PARA FIGURA" in shape.text_frame.text:
            shape.text_frame.text = ""

# Colocar figuras 2 y 4 en las dos mitades inferiores
if os.path.exists(fig2_path):
    # Figura 2 (Pie ICA): centrado en la mitad izquierda
    slide5.shapes.add_picture(fig2_path, Inches(1.85), Inches(3.20), width=Inches(3.50), height=Inches(3.50))
    print("Figura 2 agregada a Diapositiva 5")

if os.path.exists(fig4_path):
    # Figura 4 (Dispersión clima): mitad derecha
    slide5.shapes.add_picture(fig4_path, Inches(6.90), Inches(3.20), width=Inches(5.70), height=Inches(3.50))
    print("Figura 4 agregada a Diapositiva 5")

# Guardar cambios
prs.save(pptx_path)
print("\n[OK] Presentación actualizada y guardada con éxito en:", pptx_path)
