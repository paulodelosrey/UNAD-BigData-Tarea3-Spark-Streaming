from pptx import Presentation
from pptx.util import Inches
import os

script_dir = os.path.dirname(os.path.abspath(__file__))
pptx_path = os.path.join(script_dir, "Sustentación Tarea 3 Big Data UNAD.pptx")
fig5_path = os.path.join(os.path.dirname(script_dir), "output", "figura5_streaming_kafka_spark.png")

prs = Presentation(pptx_path)
slide6 = prs.slides[5]

# Remove or clear shape 0 and shape 15..19 if they were the previous placeholder card
to_remove = []
for shape in list(slide6.shapes):
    if shape.has_text_frame:
        if "OBSERVABILIDAD Y TELEMETRÍA" in shape.text_frame.text or "Auditoría en Tiempo Real" in shape.text_frame.text or ":4040" in shape.text_frame.text or "Spark UI" in shape.text_frame.text or "Inspección de Jobs" in shape.text_frame.text:
            shape.text_frame.text = ""

# Add picture
if os.path.exists(fig5_path):
    slide6.shapes.add_picture(fig5_path, Inches(0.67), Inches(3.60), width=Inches(12.0), height=Inches(3.60))
    print("[OK] Figura 5 agregada a Diapositiva 6 con éxito.")

prs.save(pptx_path)
print("[OK] Presentación guardada exitosamente.")
